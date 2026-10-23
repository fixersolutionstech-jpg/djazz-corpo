"""
ponte.py — leva os registos da Sheet (PWA, fotos, relógio) para The DJAZZ\\Corpo.

Uso:
  python ponte.py sincronizar        Sheet -> Corpo.xlsx (registo oficial) + Resumo_semana.md
  python ponte.py seco               mostra o que mudaria no Corpo.xlsx; não escreve nada
  python ponte.py resumo             só refaz o Resumo_semana.md a partir do Corpo.xlsx
  python ponte.py saude FICHEIRO.csv envia para a Sheet os dados do relógio (CSV: data,hora,item,valor,unidade,duracao_min)

Requer: pip install requests openpyxl
Lê GAS_URL e GAS_TOKEN do .env desta pasta. DJAZZ_DIR vem do .env ou do ambiente; se faltar, usa a pasta The DJAZZ do Rei.
Só este script escreve no Corpo.xlsx (fecha o Excel antes). Sem aprovação: registos apagados na PWA saem do Excel.
"""
import argparse
import csv
import hashlib
import json
import os
import shutil
import sys
import time
from datetime import date, datetime, timedelta
from pathlib import Path

AQUI = Path(__file__).parent
DJAZZ_POR_DEFEITO = r"C:\Users\Edson Adolfo\Desktop\AI\The castelo do Rei multiplicado\The DJAZZ"
META_TREINOS, META_CONVERSAS = 3, 3

# Uma folha por tipo; colunas na ordem em que aparecem no Excel
FOLHAS = {
    "treino": ("Treino", ["id", "data", "hora", "item", "series", "reps", "duracao_min", "descricao", "origem"]),
    "refeicao": ("Refeicoes", ["id", "data", "hora", "item", "descricao", "proteina_g", "kcal", "origem", "confianca"]),
    "medida": ("Medidas", ["id", "data", "hora", "item", "valor", "unidade"]),
    "social": ("Social", ["id", "data", "hora", "valor", "descricao"]),
    "saude": ("Saude", ["id", "data", "hora", "item", "valor", "unidade", "duracao_min", "origem"]),
}


def env():
    cfg = {}
    f = AQUI / ".env"
    if f.exists():
        for linha in f.read_text(encoding="utf-8").splitlines():
            linha = linha.strip()
            if linha and not linha.startswith("#") and "=" in linha:
                k, v = linha.split("=", 1)
                cfg[k.strip()] = v.strip()
    for k in ("GAS_URL", "GAS_TOKEN", "DJAZZ_DIR"):
        if os.environ.get(k):
            cfg[k] = os.environ[k]
    return cfg


def gas(cfg, corpo=None, **params):
    """Chama o GAS. Repete até 4 vezes (o Google falha de vez em quando). Nunca mostra o URL nem o token nos erros."""
    import requests

    if not cfg.get("GAS_URL") or not cfg.get("GAS_TOKEN"):
        sys.exit("Falta GAS_URL ou GAS_TOKEN no .env")
    ultimo = ""
    for i in range(4):
        try:
            if corpo is None:
                r = requests.get(cfg["GAS_URL"], params={"token": cfg["GAS_TOKEN"], **params}, timeout=60)
            else:
                r = requests.post(cfg["GAS_URL"], data=json.dumps({"token": cfg["GAS_TOKEN"], **corpo}), timeout=60)
            j = r.json()
            if j.get("erro") == "token inválido" and i < 3:
                raise ValueError("token")
            if not j.get("ok"):
                sys.exit("Erro do GAS: " + str(j.get("erro")))
            return j
        except ValueError:  # antes do RequestException: o erro de JSON do requests é das duas classes
            ultimo = "o GAS não devolveu JSON (HTTP %s)" % r.status_code
        except requests.RequestException as e:
            ultimo = "sem ligação ao GAS (%s)" % type(e).__name__  # a mensagem do requests traz o URL com o token
        time.sleep(2 * (i + 1))
    sys.exit(ultimo.capitalize() + ". Tenta outra vez daqui a pouco.")


def pasta(cfg):
    p = Path(cfg.get("DJAZZ_DIR") or DJAZZ_POR_DEFEITO) / "Corpo"
    if not p.parent.exists():
        sys.exit(f"Pasta do DJazz não encontrada: {p.parent}")
    p.mkdir(exist_ok=True)
    return p


# ---------- Excel ----------
def ler_excel(xlsx):
    """Devolve {id: registo} de todas as folhas do Corpo.xlsx (vazio se ainda não existir)."""
    from openpyxl import load_workbook

    regs = {}
    if not xlsx.exists():
        return regs
    wb = load_workbook(xlsx, read_only=True)
    for tipo, (nome, cols) in FOLHAS.items():
        if nome not in wb.sheetnames:
            continue
        for row in wb[nome].iter_rows(min_row=2, values_only=True):
            if row and row[0]:
                r = {c: ("" if v is None else v) for c, v in zip(cols, row)}
                r["tipo"] = tipo
                regs[str(r["id"])] = r
    wb.close()
    return regs


def juntar(existentes, da_sheet):
    """Une o Excel com a Sheet pelo id: a Sheet manda (correcções), os apagados saem, o que só está no Excel fica."""
    novo = dict(existentes)
    mud = {"novos": 0, "corrigidos": 0, "apagados": 0}
    for r in da_sheet:
        i = str(r["id"])
        if r.get("tipo") not in FOLHAS:
            continue
        if r.get("estado") == "apagado":
            if i in novo:
                del novo[i]
                mud["apagados"] += 1
            continue
        _, cols = FOLHAS[r["tipo"]]
        limpo = {c: r.get(c, "") for c in cols}
        limpo["tipo"] = r["tipo"]
        if i not in novo:
            mud["novos"] += 1
        elif any(str(novo[i].get(c, "")) != str(limpo[c]) for c in cols):
            mud["corrigidos"] += 1
        novo[i] = limpo
    return novo, mud


def escrever_excel(xlsx, regs):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    for tipo, (nome, cols) in FOLHAS.items():
        ws = wb.create_sheet(nome)
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for r in sorted((r for r in regs.values() if r["tipo"] == tipo), key=lambda r: (str(r["data"]), str(r["hora"]))):
            ws.append([r.get(c, "") for c in cols])
    tmp = xlsx.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    try:
        os.replace(tmp, xlsx)
    except PermissionError:
        tmp.unlink(missing_ok=True)
        sys.exit("O Corpo.xlsx está aberto. Fecha o Excel e corre outra vez.")


# ---------- resumo para o DJazz ----------
def n(v):
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def resumo(regs, hoje=None):
    """Texto do Resumo_semana.md: factos da semana (segunda a hoje) e comparação com a anterior. Sem prescrições."""
    hoje = hoje or date.today()
    ini = hoje - timedelta(days=hoje.weekday())
    ant = ini - timedelta(days=7)
    lista = list(regs.values())

    def da(a, b):
        return [r for r in lista if a.isoformat() <= str(r["data"]) < b.isoformat()]

    def semana(rs):
        treinos = [r for r in rs if r["tipo"] == "treino"]
        rel = [r for r in rs if r["tipo"] == "saude" and "treino" in str(r["item"]).lower()]
        ref = [r for r in rs if r["tipo"] == "refeicao"]
        dias_ref = len({r["data"] for r in ref})
        prot = [n(r["proteina_g"]) for r in ref if n(r["proteina_g"]) is not None]
        passos = [r for r in rs if r["tipo"] == "saude" and "passos" in str(r["item"]).lower()]
        sono = [n(r["valor"]) for r in rs if r["tipo"] == "saude" and "sono" in str(r["item"]).lower() and n(r["valor"])]
        return {
            "dias_treino": len({r["data"] for r in treinos + rel}),
            "conversas": int(sum(n(r["valor"]) or 1 for r in rs if r["tipo"] == "social")),
            "refeicoes": len(ref), "dias_ref": dias_ref,
            "prot_dia": round(sum(prot) / dias_ref) if prot and dias_ref else None,
            "passos_dia": round(sum(n(r["valor"]) or 0 for r in passos) / len({r["data"] for r in passos})) if passos else None,
            "sono_h": round(sum(sono) / len(sono), 1) if sono else None,
            "treinos": treinos,
        }

    s, a = semana(da(ini, hoje + timedelta(days=1))), semana(da(ant, ini))
    ex = {}
    for r in s["treinos"]:
        e = ex.setdefault(r["item"], {"dias": set(), "reps": 0, "min": 0})
        e["dias"].add(r["data"])
        e["reps"] += int((n(r["series"]) or 1) * (n(r["reps"]) or 0))
        e["min"] += int(n(r["duracao_min"]) or 0)

    def ultima(item):
        m = sorted((r for r in lista if r["tipo"] == "medida" and r["item"] == item), key=lambda r: (str(r["data"]), str(r["hora"])))
        return m[-1] if m else None

    def comp(x, y, suf=""):
        return f"{x}{suf}" + (f" (semana anterior: {y}{suf})" if y not in (None, 0) else "")

    L = [f"# Resumo da semana — DJazz Corpo", "",
         f"Semana de {ini:%d/%m} a {hoje:%d/%m/%Y}. Gerado pela ponte.py a {datetime.now():%d/%m %H:%M}. "
         "Estimativas de comida são aproximadas (fotos analisadas por IA ou valores indicados pelo Rei).", "",
         "| Indicador | Valor |", "|---|---|",
         f"| Dias de treino | {comp(s['dias_treino'], a['dias_treino'])} · meta {META_TREINOS} |",
         f"| Conversas novas | {comp(s['conversas'], a['conversas'])} · meta {META_CONVERSAS} |",
         f"| Refeições registadas | {s['refeicoes']} em {s['dias_ref']} dia(s) |"]
    if s["prot_dia"] is not None:
        L.append(f"| Proteína por dia registado (aprox.) | {comp(s['prot_dia'], a['prot_dia'], ' g')} |")
    if s["passos_dia"] is not None:
        L.append(f"| Passos por dia (relógio) | {comp(s['passos_dia'], a['passos_dia'])} |")
    if s["sono_h"] is not None:
        L.append(f"| Sono médio (relógio) | {comp(s['sono_h'], a['sono_h'], ' h')} |")
    for item, rotulo in (("Peso", "Último peso"), ("Cintura", "Última cintura")):
        u = ultima(item)
        if u:
            L.append(f"| {rotulo} | {u['valor']} {u['unidade']} ({u['data']}) |")
    if ex:
        L += ["", "## Treino da semana", "", "| Exercício | Dias | Reps totais | Minutos |", "|---|---|---|---|"]
        L += [f"| {k} | {len(v['dias'])} | {v['reps'] or '—'} | {v['min'] or '—'} |" for k, v in sorted(ex.items())]
    L += ["", "_Dados: `Corpo.xlsx` nesta pasta (só a ponte.py escreve nele)._", ""]
    return "\n".join(L)


# ---------- relógio (Mi Fitness -> CSV) ----------
def ler_csv_saude(ficheiro):
    """CSV com cabeçalho data,hora,item,valor,unidade,duracao_min. O id é estável: reenviar o mesmo ficheiro não duplica."""
    regs = []
    with open(ficheiro, encoding="utf-8-sig", newline="") as f:
        for l in csv.DictReader(f):
            if not l.get("data") or not l.get("item"):
                continue
            chave = "|".join(str(l.get(k, "")).strip() for k in ("data", "hora", "item"))
            regs.append({
                "id": "mf" + hashlib.sha1(chave.encode()).hexdigest()[:10],
                "data": l["data"].strip(), "hora": (l.get("hora") or "").strip()[:5], "item": l["item"].strip(),
                "valor": n(l.get("valor", "")) if n(l.get("valor", "")) is not None else "",
                "unidade": (l.get("unidade") or "").strip(),
                "duracao_min": n(l.get("duracao_min", "")) if n(l.get("duracao_min", "")) is not None else "",
                "origem": "mifitness",
            })
    return sorted(regs, key=lambda r: (r["data"], r["hora"]))


# ---------- comandos ----------
def sincronizar(cfg, seco=False):
    p = pasta(cfg)
    xlsx = p / "Corpo.xlsx"
    da_sheet = sorted(gas(cfg, acao="todos")["registos"], key=lambda r: (r["data"], r.get("hora", ""), r.get("criado_em", "")))
    existentes = ler_excel(xlsx)
    regs, mud = juntar(existentes, da_sheet)
    print(f"Sheet: {len(da_sheet)} registo(s) · Excel: {len(existentes)} -> {len(regs)} "
          f"({mud['novos']} novo(s), {mud['corrigidos']} corrigido(s), {mud['apagados']} apagado(s))")
    if seco:
        return
    if xlsx.exists() and any(mud.values()):
        bk = p / "backup"
        bk.mkdir(exist_ok=True)
        shutil.copy2(xlsx, bk / f"Corpo-{datetime.now():%Y%m%d-%H%M%S}.xlsx")
    if any(mud.values()) or not xlsx.exists():
        escrever_excel(xlsx, regs)
    (p / "Resumo_semana.md").write_text(resumo(regs), encoding="utf-8")
    print(f"Escrito: {xlsx} e Resumo_semana.md")


def main(argv):
    ap = argparse.ArgumentParser(description="Ponte DJazz Corpo: Sheet -> The DJAZZ\\Corpo")
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("sincronizar", help="Sheet -> Corpo.xlsx + Resumo_semana.md")
    sub.add_parser("seco", help="mostra o que mudaria, sem escrever")
    sub.add_parser("resumo", help="só refaz o Resumo_semana.md")
    s = sub.add_parser("saude", help="envia um CSV do relógio para a Sheet")
    s.add_argument("ficheiro")
    a = ap.parse_args(argv)
    cfg = env()
    if a.cmd == "sincronizar":
        sincronizar(cfg)
    elif a.cmd == "seco":
        sincronizar(cfg, seco=True)
    elif a.cmd == "resumo":
        p = pasta(cfg)
        (p / "Resumo_semana.md").write_text(resumo(ler_excel(p / "Corpo.xlsx")), encoding="utf-8")
        print("Resumo_semana.md refeito")
    elif a.cmd == "saude":
        regs = ler_csv_saude(a.ficheiro)
        if not regs:
            sys.exit("CSV sem linhas válidas (cabeçalho esperado: data,hora,item,valor,unidade,duracao_min)")
        for i in range(0, len(regs), 200):
            j = gas(cfg, {"acao": "saude", "regs": regs[i:i + 200]})
            print(f"Lote {i // 200 + 1}: {j['novos']} novo(s), {j['repetidos']} já existiam")


if __name__ == "__main__":
    main(sys.argv[1:])
